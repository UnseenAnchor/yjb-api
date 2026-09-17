#!/usr/bin/env python3
"""养基宝加仓记录一键刷新工具

用法:
  python refresh_yjb.py send  <手机号>     # 发送验证码
  python refresh_yjb.py pull  <验证码>     # 登录 + 拉取 + 导出 Excel
  python refresh_yjb.py export             # 用现有 token 直接导出(未过期时)
"""
import sys
import json
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from yjb_tool import YJBClient, load_token, save_token

STATE_FILE = Path(__file__).parent / "_device_id.tmp"
OUT_XLSX = Path(__file__).parent / "加仓记录_导出.xlsx"
OUT_JSON = Path(__file__).parent / "加仓记录_raw.json"
TYPE_MAP = {1: "加仓", 2: "减仓", 3: "加仓(份额)", 4: "减仓(份额)"}
STATE_MAP = {1: "待确认", 2: "已确认", 3: "已撤销", 4: "已撤销"}


def send_code(phone: str):
    device_id = __import__("uuid").uuid4().hex
    client = YJBClient(token=device_id, api="new")
    resp = client.post("/send_code", json={"phone": phone})
    STATE_FILE.write_text(json.dumps({"device_id": device_id, "phone": phone}))
    print(f"✅ 验证码已发送到 {phone[:3]}****{phone[-4:]}，响应: {resp}")


def login(code: str):
    state = json.loads(STATE_FILE.read_text())
    client = YJBClient(token=state["device_id"], api="new")
    data = client.post("/login", json={
        "phone": state["phone"], "mode": "phone", "verify_code": code,
        "invite_code": "", "is_band_wechat": 1,
    })
    token = data.get("token")
    if not token:
        print(f"❌ 登录失败: {json.dumps(data, ensure_ascii=False)[:200]}")
        sys.exit(1)
    save_token(token)
    print(f"✅ 登录成功 (昵称: {data.get('nickname', '?')})")
    return token


def with_retry(fn, tries=5, wait=8):
    """单请求级重试: 503/429/超时 自动退避重试"""
    for i in range(tries):
        try:
            return fn()
        except Exception as e:
            if i == tries - 1:
                raise
            if "503" in str(e) or "429" in str(e) or "超时" in str(e):
                time.sleep(wait * (i + 1))
            else:
                raise


def api_get(c, path, **kwargs):
    return with_retry(lambda: c.get(path, **kwargs))


def get_accounts(c):
    r = api_get(c, "/user_account")
    accounts = r.get("list", []) if isinstance(r, dict) else []
    return [(a["id"], a.get("title", "?")) for a in accounts]


def get_fund_ids(c, account_id):
    funds = []
    try:
        r = api_get(c, "/fund_hold", params={"account_id": account_id})
        for it in (r if isinstance(r, list) else []):
            if it.get("fund_id"):
                funds.append(it["fund_id"])
    except Exception:
        pass
    if not funds:  # 兜底: 自选+持仓接口
        r = api_get(c, "/position/v1/option/all")
        for it in (r if isinstance(r, list) else []):
            if it.get("is_hold") and it.get("fund_id"):
                funds.append(it["fund_id"])
    return funds


def pull_records(c, account_id, fund_ids):
    all_recs, seen = {}, set()
    for fid in fund_ids:
        page = 1
        while True:
            r = api_get(c, "/action_record", params={
                "account_id": account_id, "fund_id": fid,
                "state": 0, "type": 0, "page": page, "per_page": 100})
            batch = r if isinstance(r, list) else []
            if not batch:
                break
            for x in batch:
                if x["id"] not in seen:
                    seen.add(x["id"])
                    all_recs.setdefault(x["id"], x)
            if len(batch) < 100:
                break
            page += 1
            time.sleep(1.0)
        time.sleep(0.8)
    return list(all_recs.values())


def export(records):
    rows = []
    for x in records:
        rows.append({
            "基金ID": x.get("fund_id", ""), "基金名称": x.get("fund_name", ""),
            "类型": TYPE_MAP.get(x["type"], f"type{x['type']}"),
            "金额/份额": x.get("money", ""),
            "状态": STATE_MAP.get(x.get("state"), f"state{x.get('state')}"),
            "操作时间": x.get("operation_time", ""), "确认日": x.get("confirm_day", ""),
        })
    rows.sort(key=lambda r: r["操作时间"], reverse=True)
    OUT_JSON.write_text(json.dumps(records, ensure_ascii=False, indent=1), encoding="utf-8")

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "加仓记录"
    cols = list(rows[0].keys()) if rows else ["基金ID", "基金名称", "类型", "金额/份额", "状态", "操作时间", "确认日"]
    ws.append(cols)
    for r in rows:
        ws.append([r[c] for c in cols])
    from openpyxl.utils import get_column_letter
    for i, cname in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, len(cname) * 2 + 6)
    wb.save(OUT_XLSX)
    return rows


def pull_all(token):
    c = YJBClient(token=token, api="new")
    accounts = get_accounts(c)
    if not accounts:
        print("⚠️ 未找到持仓账户")
        return
    print("账户: " + ", ".join(f"{t}(id={i})" for i, t in accounts))
    records = []
    for acc_id, title in accounts:
        fund_ids = get_fund_ids(c, acc_id)
        print(f"  [{title}] {len(fund_ids)} 只基金，拉取记录中...")
        recs = pull_records(c, acc_id, fund_ids)
        print(f"  [{title}] {len(recs)} 条记录")
        records.extend(recs)
    rows = export(records)
    print(f"\n✅ 导出完成: {OUT_XLSX.name} ({len(rows)} 行) / {OUT_JSON.name}")
    add = sum(1 for r in rows if r["类型"].startswith("加仓"))
    dec = sum(1 for r in rows if r["类型"].startswith("减仓"))
    print(f"   加仓 {add} 笔 / 减仓 {dec} 笔")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    cmd = sys.argv[1]
    if cmd == "send" and len(sys.argv) >= 3:
        send_code(sys.argv[2])
    elif cmd == "pull" and len(sys.argv) >= 3:
        token = login(sys.argv[2])
        pull_all(token)
    elif cmd == "export":
        token = load_token()
        if not token:
            print("❌ 无保存的 token，请先 send + pull")
            sys.exit(1)
        try:
            pull_all(token)
        except Exception as e:
            print(f"❌ token 可能已过期: {e}")
            print("请运行: python refresh_yjb.py send <手机号>")
            sys.exit(1)
    else:
        print(__doc__)
